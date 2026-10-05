import os
import io
import json
import streamlit as st
import pandas as pd
import plotly.express as px
from datetime import datetime
from docx import Document

# Thư viện kết nối Google Sheets
import gspread
from google.oauth2.service_account import Credentials

st.set_page_config(page_title="Hệ thống Quản lý & Báo cáo KH&CN", page_icon="📊", layout="wide")

GOOGLE_SHEET_NAME = "QuanLy_DonDangKy_KHCN" 

# --- 1. KHỞI TẠO KẾT NỐI GOOGLE SHEETS ---
def get_gspread_client():
    scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
    
    # Ưu tiên kiểm tra file local google_key.json trước
    if os.path.exists("google_key.json"):
        try:
            credentials = Credentials.from_service_account_file("google_key.json", scopes=scopes)
            return gspread.authorize(credentials)
        except Exception as e:
            st.error(f"Lỗi đọc file google_key.json: {e}")
            return None
            
    # Đọc từ Secrets nếu chạy trên Streamlit Cloud
    try:
        if hasattr(st, "secrets") and "gcp_service_account" in st.secrets:
            service_account_info = json.loads(st.secrets["gcp_service_account"]["json_key"])
            credentials = Credentials.from_service_account_info(service_account_info, scopes=scopes)
            return gspread.authorize(credentials)
    except Exception:
        pass
        
    return None

@st.cache_data(ttl=10)
def load_data_from_gsheet():
    try:
        gc = get_gspread_client()
        if not gc:
            return pd.DataFrame()
            
        sh = gc.open(GOOGLE_SHEET_NAME)
        ws = sh.worksheet("QUAN_LY_HO_SO")
        
        data = ws.get_all_values()
        if len(data) <= 1:
            return pd.DataFrame()
            
        headers = data[1] # Hàng 2 là tiêu đề cột
        rows = data[2:]   # Dữ liệu từ hàng 3
        
        # Xử lý tên cột trùng/trống
        clean_headers = []
        counts = {}
        for idx, h in enumerate(headers):
            name = str(h).strip()
            if not name:
                name = f"Unnamed_{idx}"
            if name in counts:
                counts[name] += 1
                name = f"{name}_{counts[name]}"
            else:
                counts[name] = 0
            clean_headers.append(name)
        
        df = pd.DataFrame(rows, columns=clean_headers)
        
        # Ép kiểu dữ liệu số
        num_cols = ["STT", "Tổng kinh phí đề xuất (VNĐ)", "Kinh phí NSNN (VNĐ)", "Kinh phí Ngoài NSNN (VNĐ)", "Thời gian thực hiện (Tháng)"]
        for col in num_cols:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col].astype(str).str.replace(',', ''), errors='coerce').fillna(0)
                
        return df
    except Exception as e:
        st.error(f"Lỗi khi tải dữ liệu từ Google Sheets: {e}")
        return pd.DataFrame()

# --- 2. HÀM TẠO FILE WORD BÁO CÁO TỰ ĐỘNG ---
def generate_docx_report(df):
    doc = Document()
    doc.add_heading('BÁO CÁO TỔNG HỢP HỒ SƠ KH&CN', 0)
    
    doc.add_paragraph(f"Thời gian xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    doc.add_paragraph(f"Tổng số lượng hồ sơ: {len(df)}")
    
    # Thống kê ngân sách
    tong_kp = df["Tổng kinh phí đề xuất (VNĐ)"].sum() if "Tổng kinh phí đề xuất (VNĐ)" in df.columns else 0
    kp_nsnn = df["Kinh phí NSNN (VNĐ)"].sum() if "Kinh phí NSNN (VNĐ)" in df.columns else 0
    
    doc.add_heading('1. Thống kê Kinh phí', level=1)
    doc.add_paragraph(f"- Tổng kinh phí đề xuất: {tong_kp:,.0f} VNĐ")
    doc.add_paragraph(f"- Tổng kinh phí NSNN: {kp_nsnn:,.0f} VNĐ")
    
    doc.add_heading('2. Danh sách hồ sơ chi tiết', level=1)
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'STT'
    hdr_cells[1].text = 'Tên nhiệm vụ'
    hdr_cells[2].text = 'Lĩnh vực'
    hdr_cells[3].text = 'Chủ trì'
    
    for idx, row in df.iterrows():
        row_cells = table.add_row().cells
        row_cells[0].text = str(row.get('STT', idx + 1))
        row_cells[1].text = str(row.get('Tên nhiệm vụ / Cụm / Chuỗi', ''))
        row_cells[2].text = str(row.get('Lĩnh vực', ''))
        row_cells[3].text = str(row.get('Tên tổ chức chủ trì', ''))
        
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# --- 3. GIAO DIỆN VÀ CHỨC NĂNG ---
st.title("📌 HỆ THỐNG QUẢN LÝ & BÁO CÁO HỒ SƠ KH&CN (GOOGLE SHEETS)")

df_raw = load_data_from_gsheet()

if df_raw.empty:
    st.warning("⚠️ Không có dữ liệu hoặc không thể kết nối đến Google Sheets. Vui lòng kiểm tra lại cấu hình.")
else:
    # --- THANH BỘ LỌC BÊN TRÁI (SIDEBAR) ---
    st.sidebar.header("🔍 Bộ lọc tìm kiếm")
    
    linh_vuc_list = ["Tất cả"] + list(df_raw["Lĩnh vực"].dropna().unique()) if "Lĩnh vực" in df_raw.columns else ["Tất cả"]
    selected_linh_vuc = st.sidebar.selectbox("Chọn Lĩnh vực:", linh_vuc_list)
    
    trang_thai_col = [c for c in df_raw.columns if "Trạng thái" in c]
    trang_thai_name = trang_thai_col[0] if trang_thai_col else None
    
    if trang_thai_name:
        trang_thai_list = ["Tất cả"] + list(df_raw[trang_thai_name].dropna().unique())
        selected_trang_thai = st.sidebar.selectbox("Chọn Trạng thái:", trang_thai_list)
    else:
        selected_trang_thai = "Tất cả"
        
    # Áp dụng bộ lọc
    df_filtered = df_raw.copy()
    if selected_linh_vuc != "Tất cả":
        df_filtered = df_filtered[df_filtered["Lĩnh vực"] == selected_linh_vuc]
    if trang_thai_name and selected_trang_thai != "Tất cả":
        df_filtered = df_filtered[df_filtered[trang_thai_name] == selected_trang_thai]

    # --- METRICS BÁO CÁO NHANH ---
    col1, col2, col3 = st.columns(3)
    col1.metric("Tổng số hồ sơ", len(df_filtered))
    
    tong_kp = df_filtered["Tổng kinh phí đề xuất (VNĐ)"].sum() if "Tổng kinh phí đề xuất (VNĐ)" in df_filtered.columns else 0
    kp_ns = df_filtered["Kinh phí NSNN (VNĐ)"].sum() if "Kinh phí NSNN (VNĐ)" in df_filtered.columns else 0
    
    col2.metric("Tổng kinh phí", f"{tong_kp:,.0f} VNĐ")
    col3.metric("Kinh phí NSNN", f"{kp_ns:,.0f} VNĐ")

    # --- TAB CHỨC NĂNG ---
    tab1, tab2, tab3 = st.tabs(["📋 Danh sách Hồ sơ", "📊 Biểu đồ Thống kê", "📄 Xuất Báo cáo Word"])

    with tab1:
        st.subheader("Bảng dữ liệu hồ sơ chi tiết")
        st.dataframe(df_filtered, use_container_width=True)

    with tab2:
        st.subheader("Phân tích & Thống kê trực quan")
        c1, c2 = st.columns(2)
        
        with c1:
            if "Lĩnh vực" in df_filtered.columns:
                fig1 = px.pie(df_filtered, names="Lĩnh vực", title="Tỷ lệ hồ sơ theo Lĩnh vực", hole=0.4)
                st.plotly_chart(fig1, use_container_width=True)
                
        with c2:
            if trang_thai_name:
                fig2 = px.bar(df_filtered, x=trang_thai_name, title="Phân bố theo Trạng thái hồ sơ", color=trang_thai_name)
                st.plotly_chart(fig2, use_container_width=True)

    with tab3:
        st.subheader("Tự động xuất báo cáo theo biểu mẫu Word")
        st.write("Bấm vào nút bên dưới để hệ thống tự động tổng hợp danh sách và số liệu kinh phí ra file `.docx`:")
        
        docx_bytes = generate_docx_report(df_filtered)
        st.download_button(
            label="📥 Tải Báo cáo Word (.docx)",
            data=docx_bytes,
            file_name=f"Bao_Cao_KHCN_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )